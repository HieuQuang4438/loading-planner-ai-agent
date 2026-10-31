"""How each kind of file is handed to the model. A spreadsheet is read here, on the server, into lines of text that name
their row ("Packing List · dòng 6"), so the place the model reports is a real one; a PDF or an image is sent as is."""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from pathlib import Path

MIME = {".pdf": "application/pdf", ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}
SHEETS = {".xlsx", ".xlsm", ".csv"}
MAX_ROWS = 2000  # a sheet longer than this is cut, and the cut is said in the text


@dataclass(frozen=True)
class Upload:
    """One file of a lot: what the model gets, either `text` (a spreadsheet read here) or `data` with its `mime`."""

    name: str
    text: str | None = None
    data: bytes | None = None
    mime: str | None = None


class UnsupportedFile(ValueError):
    pass


def load(name: str, data: bytes) -> Upload:
    suffix = Path(name).suffix.lower()
    if suffix in SHEETS:
        return Upload(name, text=sheet_text(name, data))
    if suffix in MIME:
        return Upload(name, data=data, mime=MIME[suffix])
    raise UnsupportedFile(f"{name}: chỉ nhận Excel, CSV, PDF, ảnh")


def sheet_text(name: str, data: bytes) -> str:
    """Every non-empty row as `[<sheet> · dòng <n>] A=... | B=...`, cell values as shown (`data_only`)."""
    suffix = Path(name).suffix.lower()
    out: list[str] = []
    if suffix == ".csv":
        text = data.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
        out += _rows("CSV", rows)
    else:
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for sheet in book.worksheets:
            out += _rows(sheet.title, [[c for c in row] for row in sheet.iter_rows(values_only=True)])
    return "\n".join(out)


def _rows(sheet: str, rows: list[list[object]]) -> list[str]:
    out = []
    for n, row in enumerate(rows, 1):
        if n > MAX_ROWS:
            out.append(f"[{sheet}] … còn {len(rows) - MAX_ROWS} dòng không đọc")
            break
        cells = [f"{_column(i)}={_cell(v)}" for i, v in enumerate(row) if v is not None and str(v).strip()]
        if cells:
            out.append(f"[{sheet} · dòng {n}] " + " | ".join(cells))
    return out


def _cell(value: object) -> str:
    return " ".join(str(value).split())


def _column(index: int) -> str:
    name = ""
    index += 1
    while index:
        index, rest = divmod(index - 1, 26)
        name = chr(65 + rest) + name
    return name
